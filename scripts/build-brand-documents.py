from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak,Image
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import json

SITE=Path(__file__).resolve().parents[1];OUT=SITE/'public/brand-assets/v1';DOC=OUT/'documentos';DOC.mkdir(exist_ok=True)
PAPER='F5F2ED';INK='171513';ORANGE='FA6404';MUTED='66615B'
LOGO=OUT/'logos/lgh-assinatura-principal-claro.png'

def cellshade(cell,color):
 p=cell._tc.get_or_add_tcPr();s=OxmlElement('w:shd');s.set(qn('w:fill'),color);p.append(s)
def document(title):
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(2.1);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2.15);s.right_margin=Cm(2.15)
 s.header_distance=Cm(.65);s.footer_distance=Cm(.7)
 for name in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
  st=d.styles[name];st.font.name='Manrope' if name in ['Normal','Body Text','Subtitle'] else 'Space Grotesk SemiBold';st.font.color.rgb=RGBColor.from_string('000000');st.font.size=Pt(10.5)
  rf=st.element.get_or_add_rPr().rFonts
  for attr in list(rf.attrib):
   if 'theme' in attr.lower():del rf.attrib[attr]
  st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.14
 d.styles['Title'].font.size=Pt(27);d.styles['Title'].font.bold=False;d.styles['Title'].paragraph_format.space_after=Pt(14)
 for border in d.styles.element.xpath('.//w:pBdr/*'):
  border.set(qn('w:color'),ORANGE)
  for attr in ['themeColor','themeTint','themeShade']:
   if qn('w:'+attr) in border.attrib:del border.attrib[qn('w:'+attr)]
 for n,size in [('Heading 1',16),('Heading 2',12)]:d.styles[n].font.size=Pt(size);d.styles[n].font.bold=False;d.styles[n].paragraph_format.space_before=Pt(16)
 hp=s.header.paragraphs[0];hp.add_run().add_picture(str(LOGO),width=Cm(5.2))
 for inline in hp._p.xpath('.//wp:docPr'):inline.set('descr','Lucas Gil Films')
 fp=s.footer.paragraphs[0];fp.paragraph_format.space_before=Pt(0)
 r=fp.add_run('@lghfilms  •  +55 11 91511-7067');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string(MUTED)
 fp.add_run(' '*8);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
 d.core_properties.author='Lucas Gil Films';d.core_properties.title=title;d.core_properties.subject='Modelo editável';d.core_properties.keywords='Lucas Gil Films, LGH, modelo'
 return d
def para(d,text,style=None):return d.add_paragraph(text,style)
def title(d,text):d.add_paragraph(text,'Title')
def section(d,text):d.add_paragraph(text,'Heading 1')
def field(d,label,prompt):
 p=d.add_paragraph();r=p.add_run(label+'  ');r.bold=True;r=p.add_run('['+prompt+']');r.font.color.rgb=RGBColor.from_string(MUTED)
 return p
def table(d,headers,rows,widths):
 t=d.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Cm(w)
 for i,h in enumerate(headers):
  c=t.rows[0].cells[i];c.width=Cm(widths[i]);cellshade(c,INK);r=c.paragraphs[0].add_run(h);r.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(9)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for j,row in enumerate(rows):
  cells=t.add_row().cells
  for i,val in enumerate(row):
   c=cells[i];c.width=Cm(widths[i]);c.paragraphs[0].add_run(str(val));cellshade(c,PAPER if j%2==0 else 'FFFFFF')
 for row in t.rows:
  for c in row.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b);mar=OxmlElement('w:tcMar')
   for e in ['top','left','bottom','right']:
    x=OxmlElement('w:'+e);x.set(qn('w:w'),'95');x.set(qn('w:type'),'dxa');mar.append(x)
   pr.append(mar)
   for p in c.paragraphs:p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.12
 return t

d=document('Proposta de produção audiovisual');title(d,'Proposta de produção audiovisual')
para(d,'Produção e edição de vídeos para [nome do cliente]. Esta proposta apresenta o objetivo, as entregas, o calendário e o investimento do projeto.')
for l,p in [('Cliente','empresa e pessoa de contato'),('Projeto','nome do projeto'),('Data e validade','data da proposta e prazo de validade')]:field(d,l,p)
section(d,'Objetivo do projeto');para(d,'[Descreva a mensagem do filme, o público e o resultado esperado. Use um parágrafo curto que explique o que o cliente precisa comunicar.]')
section(d,'Entregas previstas')
table(d,['Entrega','Quantidade','Formato e duração'],[['Filme principal','[qtd.]','[formato / duração]'],['Versões para redes','[qtd.]','[proporção / duração]'],['Legendas e idiomas','[qtd.]','[arquivo / idioma]'],['Outros materiais','[qtd.]','[descrever]']],[6.1,2.2,8.4])
section(d,'Direção criativa');para(d,'[Descreva a linguagem, as referências visuais e sonoras e os elementos essenciais da narrativa.]')
d.add_page_break();title(d,'Produção e calendário')
para(d,'As datas e os responsáveis abaixo serão confirmados com o cliente antes do início da produção.')
table(d,['Etapa','Resultado esperado','Prazo'],[['Briefing','Objetivo e referências aprovados','[data]'],['Pré-produção','Roteiro e plano de captação','[data]'],['Captação','Material previsto no roteiro','[data]'],['Edição','Primeira versão do filme','[data]'],['Revisões','Ajustes consolidados pelo cliente','[data]'],['Entrega','Arquivos finais aprovados','[data]']],[3.7,9.3,3.7])
section(d,'Revisões e aprovações')
for l,p in [('Rodadas incluídas','quantidade de rodadas de revisão'),('Responsável pela aprovação','nome e contato'),('Prazo para retorno','prazo combinado para enviar comentários')]:field(d,l,p)
para(d,'Os comentários serão reunidos em uma devolutiva por rodada. Mudanças nas entregas ou no roteiro aprovado terão seu impacto em prazo e valor alinhado antes da execução.')
section(d,'Materiais e logística')
field(d,'Materiais do cliente','logos, referências, textos e acessos necessários');field(d,'Captação','local, data, duração e participantes');field(d,'Equipe e equipamentos','recursos previstos na proposta')
d.add_page_break();title(d,'Investimento e condições')
table(d,['Item','Valor'],[['Pré-produção','R$ [valor]'],['Captação','R$ [valor]'],['Edição e finalização','R$ [valor]'],['Despesas previstas','R$ [valor]'],['Investimento total','R$ [valor]']],[12.2,4.5])
section(d,'Condições comerciais')
for l,p in [('Pagamento','parcelas, valores e vencimentos'),('Despesas','deslocamento, locação e outros itens incluídos ou cobrados à parte'),('Reagendamento','condições a combinar'),('Direitos de uso','canais, território, período e licenças necessárias'),('Arquivos originais','condições de entrega de brutos e arquivos editáveis'),('Armazenamento','prazo de disponibilidade dos arquivos finais')]:field(d,l,p)
section(d,'Aprovação')
para(d,'[Nome do cliente e responsável] confirma o escopo, o investimento e as condições desta proposta em [data].')
field(d,'Responsável pelo projeto','Lucas Gil Henriques');field(d,'Registro da aprovação','assinatura ou referência da confirmação por escrito')
d.save(DOC/'lgh-proposta-comercial.docx')

d=document('Briefing audiovisual');title(d,'Briefing audiovisual')
para(d,'Preencha este briefing para alinhar a produção do vídeo. As respostas orientam a proposta, o roteiro e o plano de captação.')
for l,p in [('Cliente e projeto','preencher'),('Responsável e contato','preencher')]:field(d,l,p)
for heading,prompt in [('Objetivo','O que o vídeo precisa comunicar e qual ação o público deve tomar?'),('Público','Quem vai assistir e o que essas pessoas já sabem sobre o assunto?'),('Mensagem principal','Qual é a ideia que deve permanecer depois do filme?'),('Canais e formatos','Onde o vídeo será exibido? Informe proporções, duração e idiomas.'),('Referências','Inclua links e explique o que gosta em cada referência.')]:
 section(d,heading);para(d,prompt);para(d,'[Resposta]')
d.add_page_break();title(d,'Produção e entrega')
for heading,prompt in [('Conteúdo e participantes','Quem aparece no vídeo? Quais cenas, falas ou produtos são essenciais?'),('Local e recursos','Onde será a captação? Há restrições de acesso, som, luz ou horário?'),('Materiais disponíveis','Liste logos, imagens, roteiros, trilhas e outros materiais que podem ser utilizados.'),('Datas e investimento','Informe data de captação, prazo de entrega e faixa de investimento prevista.'),('Aprovação','Quem aprova o roteiro e a edição? Qual é o prazo para retorno?'),('Observações','Registre autorizações, necessidades de acessibilidade e pontos que exigem alinhamento.')]:
 section(d,heading);para(d,prompt);para(d,'[Resposta]')
d.save(DOC/'lgh-briefing-audiovisual.docx')

d=document('Papel timbrado Lucas Gil Films');title(d,'[Título do documento]')
field(d,'Destinatário','nome ou empresa');field(d,'Data','dia, mês e ano')
para(d,'[Escreva a mensagem aqui. Apresente o assunto e o motivo do contato no primeiro parágrafo.]')
para(d,'[Desenvolva as informações, o pedido ou o encaminhamento necessário.]')
para(d,'[Indique o próximo passo e o prazo, quando houver.]')
para(d,'Lucas Gil Henriques\nLucas Gil Films\n+55 11 91511-7067')
d.save(DOC/'lgh-papel-timbrado.docx')

# Short standalone guide, also downloadable from the local brand book.
pdfmetrics.registerFont(TTFont('LGHDisplay',str(OUT/'fontes/SpaceGrotesk-Semibold.ttf')))
pdfmetrics.registerFont(TTFont('LGHBody',str(OUT/'fontes/Manrope-Regular.ttf')))
styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='LGHTitle',fontName='LGHDisplay',fontSize=28,leading=33,textColor=colors.HexColor('#171513'),spaceAfter=18))
styles.add(ParagraphStyle(name='LGHHead',fontName='LGHDisplay',fontSize=16,leading=21,textColor=colors.HexColor('#171513'),spaceBefore=16,spaceAfter=8))
styles.add(ParagraphStyle(name='LGHBody',fontName='LGHBody',fontSize=10,leading=15,textColor=colors.HexColor('#171513'),spaceAfter=9))
story=[]
def p(txt,style='LGHBody'):story.append(Paragraph(txt,styles[style]))
def head(txt):p(txt,'LGHHead')
def newpage():story.append(PageBreak())
story.append(Image(str(LOGO),width=260,height=65));story.append(Spacer(1,16));p('Guia do kit de marca','LGHTitle')
p('Arquivos e modelos para a comunicação de Lucas Gil Films. Escolha a peça pelo uso, preserve as proporções e preencha os campos dos templates antes de compartilhar.')
head('A assinatura')
p('A marca principal é <b>lucas gil films.</b>, em minúsculas. O monograma <b>lgh.</b> atende aos espaços menores. O ponto faz parte da assinatura.')
head('Quatro versões de logo')
for txt in ['Principal para fundo claro: letras #171513 e ponto #FA6404.','Principal para fundo escuro: letras #F5F2ED e ponto #FA6404.','Monocromática preta: letras e ponto #000000.','Monocromática branca: letras e ponto #FFFFFF.']:p(txt)
head('Avatares')
p('Use o avatar preto sobre laranja como versão principal. A alternativa clara usa letras escuras e ponto laranja. Há também uma versão escura. Os arquivos quadrados reservam margem para o recorte circular.')
head('Área de proteção')
p('Deixe ao redor do logo pelo menos a altura da letra l. Preserve a relação entre letras e ponto. Em espaços pequenos, prefira o monograma. Confira a leitura no tamanho final, especialmente nos ícones de 16 e 32 px.')
newpage();p('Arquivos e cores','LGHTitle')
for h,t in [('SVG','Arquivo vetorial para uso digital. Os logos têm letras convertidas em contornos. Os templates de redes e vídeo mantêm os textos editáveis e precisam das fontes do kit instaladas.'),('PNG','Imagem pronta para usar. Os logos, marcas d’água e identificações têm transparência. Avatares, capas e cartelas têm fundo. O texto de um PNG não é editável.'),('PDF vetorial','Formato para encaminhar a gráficas e fornecedores. Os logos mantêm os contornos vetoriais. Os PDFs do kit usam cores RGB; combine a conversão para o perfil CMYK e o suporte de impressão com a gráfica.'),('Fontes','Space Grotesk Semibold para títulos e Manrope Regular para texto. Instale os TTF da pasta fontes para editar os modelos Office e SVG. As fontes acompanham suas licenças SIL OFL 1.1.')]:head(h);p(t)
head('Cores de apoio');p('Laranja #FA6404, coral #E2725B, rosa #D58D8D, amarelo #FFDD44. Tinta #171513 e papel #F5F2ED formam a base. O laranja profundo #C44700 pode receber texto branco. Use as cores de apoio nos materiais, preservando as versões oficiais do logo.')
newpage();p('Capas e conteúdo','LGHTitle')
for h,t in [('LinkedIn pessoal','1584 × 396 px. A composição reserva o lado esquerdo para a sobreposição da foto de perfil.'),('LinkedIn página','4200 × 700 px. Arquivos PNG e JPG com composição própria para a página da empresa.'),('Facebook','1702 × 630 px, matriz em resolução dupla de 851 × 315. A exibição pode recortar a capa de formas diferentes entre dispositivos. Confira a prévia antes de salvar.'),('YouTube','2560 × 1440 px. A assinatura e a descrição ficam dentro da área central de aproximadamente 1544 × 422 px, preservando a leitura nos diferentes recortes.'),('Posts e vídeo','Modelos quadrados 1080 × 1080, feed 1080 × 1350, stories 1080 × 1920 e thumbnails 1280 × 720. São matrizes de produção, não uma lista de todos os limites de cada plataforma.')]:head(h);p(t)
p('Os avatares específicos são exportações de uso do kit. Os tamanhos não representam mínimos oficiais das plataformas. Reveja o recorte circular no upload.')
head('Referências de dimensões')
for url in ['https://www.linkedin.com/help/linkedin/answer/a568217','https://www.linkedin.com/help/linkedin/answer/a563309','https://support.google.com/youtube/answer/10456525','https://www.facebook.com/help/125379114252045']:
 p(f'<link href="{url}" color="#C44700">{url}</link>')
p('Referências consultadas em 9 de setembro de 2026. A página da Meta exigiu autenticação nesta consulta; a capa do Facebook é uma matriz de produção a conferir no upload.')
newpage();p('Modelos para trabalhar','LGHTitle')
for h,t in [('Proposta comercial','Documento editável com objetivo, escopo, etapas, prazos, investimento e condições. Substitua todos os campos entre colchetes. Os valores e os termos devem ser definidos para cada projeto.'),('Briefing','Questionário de duas páginas para iniciar o projeto. Preencha com o cliente e use as respostas para montar a proposta e o roteiro.'),('Apresentação','Modelo 16:9 adaptado da estrutura do template Aotta, com a marca LGH. Use o POTX para abrir uma nova apresentação ou duplique os slides do PPTX. Há capas, seções, títulos de uma e duas linhas, uma e duas colunas, imagem, cronograma e investimento.'),('Cronograma','Defina a data inicial e ajuste os dias e durações de cada etapa. Datas, barras semanais e totais acompanham as entradas. As durações são exemplos de planejamento e usam dias corridos.'),('Lista de planos','Registre cenas, locais, enquadramentos, ações, áudio e arquivos. Atualize a situação de cada plano e acompanhe a contagem de itens gravados.'),('Papelaria','Papel timbrado em Word, assinatura de e-mail em HTML e PNG, cartão digital e contato VCF. A assinatura HTML usa texto e links sem depender de uma imagem hospedada.')]:head(h);p(t)
def footer(canvas,doc):
 canvas.setFont('LGHBody',8);canvas.setFillColor(colors.HexColor('#66615B'));canvas.drawString(48,26,'Lucas Gil Films  |  Kit de marca v1');canvas.drawRightString(547,26,str(doc.page))
SimpleDocTemplate(str(DOC/'lgh-guia-do-kit.pdf'),pagesize=(595.276,841.89),rightMargin=48,leftMargin=48,topMargin=40,bottomMargin=44,title='Guia do kit de marca Lucas Gil Films',author='Lucas Gil Films').build(story,onFirstPage=footer,onLaterPages=footer)
print('Created proposal, briefing, letterhead and PDF guide.')
